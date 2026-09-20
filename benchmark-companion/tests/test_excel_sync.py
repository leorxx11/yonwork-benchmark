from __future__ import annotations

import tempfile
import unittest
import uuid
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from benchmark_companion.excel_sync import (
    ALL_HEADERS,
    EXCEL_DATETIME_FORMAT,
    EXCEL_INTEGER_FORMAT,
    ExcelResultsSync,
)
from benchmark_companion.models import (
    RunRecord,
    RunStatus,
    SyncStatus,
    TokenStatus,
)


class ExcelSyncTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp())
        self.path = self.temp_dir / "benchmark.xlsx"
        workbook = Workbook()
        results = workbook.active
        results.title = "Results"
        results.append(list(ALL_HEADERS[:14]))
        results.append([None] * 14)
        results.append(
            [
                "YonWork",
                "Existing",
                1,
                "old",
                datetime(2026, 1, 1, 10, 0),
                datetime(2026, 1, 1, 10, 0, 1),
                1.0,
                "deepseek-flash",
                1,
                0,
                10,
                2,
                12,
                1,
            ]
        )
        table = Table(displayName="ResultsTable", ref="A1:N3")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
        )
        results.add_table(table)
        cases = workbook.create_sheet("Cases")
        cases.append(["CaseName", "Prompt", "Runs", "Enabled"])
        cases.append(["Case01", "你好", 1, True])
        workbook.save(self.path)
        workbook.close()
        self.backups = self.temp_dir / "backups"

    def tearDown(self) -> None:
        for path in sorted(self.temp_dir.rglob("*"), reverse=True):
            if path.is_file():
                path.unlink(missing_ok=True)
            elif path.is_dir():
                path.rmdir()
        self.temp_dir.rmdir()

    def _record(self, *, token_status: TokenStatus = TokenStatus.SUCCESS) -> RunRecord:
        record_id = str(uuid.uuid4())
        return RunRecord(
            record_id=record_id,
            session_id=str(uuid.uuid4()),
            position=0,
            product="WorkBuddy",
            case_name="Case01",
            run_no=1,
            prompt="你好",
            start_time="2026-09-14T10:00:00.000+08:00",
            end_time="2026-09-14T10:00:02.500+08:00",
            duration_seconds=2.5,
            model="deepseek-flash",
            api_calls=2,
            error_calls=0,
            input_tokens=82676,
            output_tokens=1237,
            total_tokens=83913,
            api_use_time=2,
            status=RunStatus.SUCCESS,
            token_status=token_status,
            note="",
            sync_status=SyncStatus.PENDING,
            workbook_path=str(self.path),
            stats_script_path="C:/stats.ps1",
            token_name="workbuddy",
            created_at="2026-09-14T10:00:02.500+08:00",
            updated_at="2026-09-14T10:00:02.500+08:00",
        )

    def test_appends_extends_table_and_deduplicates(self) -> None:
        record = self._record()
        sync = ExcelResultsSync(self.path, self.backups)
        first = sync.sync_records([record])
        first_row = first.row_by_record[record.record_id]
        second = sync.sync_records([record])
        self.assertEqual(first_row, second.row_by_record[record.record_id])

        workbook = load_workbook(self.path, read_only=False, data_only=False)
        try:
            sheet = workbook["Results"]
            self.assertEqual(ALL_HEADERS, tuple(sheet.cell(1, col).value for col in range(1, 19)))
            self.assertEqual(record.record_id, sheet.cell(first_row, 18).value)
            self.assertEqual("A1:R4", next(iter(sheet.tables.values())).ref)
            self.assertGreaterEqual(sheet.column_dimensions["P"].width, 16)
            self.assertGreaterEqual(sheet.column_dimensions["R"].width, 40)
            self.assertIsInstance(sheet.cell(first_row, 5).value, datetime)
            self.assertIsInstance(sheet.cell(first_row, 6).value, datetime)
            self.assertEqual(EXCEL_DATETIME_FORMAT, sheet.cell(first_row, 5).number_format)
            self.assertEqual(EXCEL_DATETIME_FORMAT, sheet.cell(first_row, 6).number_format)
            for column in range(9, 14):
                self.assertEqual(EXCEL_INTEGER_FORMAT, sheet.cell(first_row, column).number_format)
                self.assertIsInstance(sheet.cell(first_row, column).value, int)
            ids = [sheet.cell(row, 18).value for row in range(2, sheet.max_row + 1)]
            self.assertEqual(1, ids.count(record.record_id))
        finally:
            workbook.close()

    def test_no_data_writes_blank_token_values(self) -> None:
        record = self._record(token_status=TokenStatus.NO_DATA)
        outcome = ExcelResultsSync(self.path, self.backups).sync_records([record])
        row = outcome.row_by_record[record.record_id]
        workbook = load_workbook(self.path, read_only=True, data_only=False)
        try:
            sheet = workbook["Results"]
            self.assertIsNone(sheet.cell(row, 9).value)
            self.assertEqual("NoData", sheet.cell(row, 16).value)
        finally:
            workbook.close()

    def test_plain_results_range_does_not_require_native_table(self) -> None:
        workbook = load_workbook(self.path, read_only=False, data_only=False)
        sheet = workbook["Results"]
        del sheet.tables["ResultsTable"]
        workbook.save(self.path)
        workbook.close()

        record = self._record()
        outcome = ExcelResultsSync(self.path, self.backups).sync_records([record])
        workbook = load_workbook(self.path, read_only=False, data_only=False)
        try:
            sheet = workbook["Results"]
            row = outcome.row_by_record[record.record_id]
            self.assertEqual(record.record_id, sheet.cell(row, 18).value)
            self.assertEqual(0, len(sheet.tables))
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
