from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("BENCHMARK_COMPANION_DISABLE_GLOBAL_HOTKEYS", "1")

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox

from benchmark_companion.config import AppConfig
from benchmark_companion.excel_sync import BASE_HEADERS
from benchmark_companion.fonts import load_ui_font
from benchmark_companion.models import ModelMode, RunStatus
from benchmark_companion.storage import BenchmarkStore
from benchmark_companion.ui import MainWindow


class UiFlowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.application = QApplication.instance() or QApplication([])
        cls.application.setFont(load_ui_font(10))

    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.workbook_path = self.root / "benchmark.xlsx"
        workbook = Workbook()
        results = workbook.active
        results.title = "Results"
        results.append(list(BASE_HEADERS))
        results.append([None] * len(BASE_HEADERS))
        table = Table(displayName="ResultsTable", ref="A1:N2")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
        )
        results.add_table(table)
        cases = workbook.create_sheet("Cases")
        cases.append(["CaseName", "Prompt", "Runs", "Enabled"])
        cases.append(["Case01", "你好！", 1, True])
        long_text = workbook.create_sheet("long-text")
        long_text.append(["CaseName", "Prompt", "Runs", "Enabled"])
        long_text.append(["Long01", "这是长文本提示词", 1, True])
        workbook.save(self.workbook_path)
        workbook.close()

        self.script_path = self.root / "stats.ps1"
        self.script_path.write_text('Write-Output "1|0|10|2|12|1"', encoding="utf-8")
        self.config = AppConfig(
            workbook_path=str(self.workbook_path),
            stats_script_path=str(self.script_path),
            token_name="workbuddy",
            database_path=str(self.root / "data" / "companion.db"),
            backup_dir=str(self.root / "data" / "backups"),
            log_path=str(self.root / "data" / "app.log"),
            always_on_top=False,
        )
        self.config.save(self.root / "config.json")
        self.store = BenchmarkStore(self.config.database)
        self.window = MainWindow(self.config, self.store)
        self.window.show()
        QTest.qWait(50)

    def tearDown(self) -> None:
        for worker in list(self.window._workers):
            worker.wait(30_000)
        self.application.processEvents()
        self.window.close()
        self.application.processEvents()
        self.temporary.cleanup()

    def test_default_model_full_flow_persists_and_syncs(self) -> None:
        index = self.window.mode_combo.findData(ModelMode.WORKBUDDY_DEFAULT.value)
        self.window.mode_combo.setCurrentIndex(index)
        self.window.start_new_batch()
        self.assertEqual("Case01", self.window.current_item.case_name)

        self.window.start_run()
        QTest.qWait(25)
        self.window.finish_run(RunStatus.SUCCESS)

        worker = self.window._sync_worker
        self.assertIsNotNone(worker)
        self.assertTrue(worker.wait(10_000))
        QTest.qWait(50)
        self.assertIsNone(self.window.current_item)
        self.assertEqual(0, self.store.count_pending(), self.window.status_label.text())

        workbook = load_workbook(self.workbook_path, read_only=True, data_only=False)
        try:
            sheet = workbook["Results"]
            self.assertEqual("WorkBuddy", sheet.cell(2, 1).value)
            self.assertEqual("Success", sheet.cell(2, 15).value)
            self.assertEqual("NotApplicable", sheet.cell(2, 16).value)
        finally:
            workbook.close()

    def test_new_batch_can_switch_model_without_skipping_remaining_items(self) -> None:
        default_index = self.window.mode_combo.findData(
            ModelMode.WORKBUDDY_DEFAULT.value
        )
        self.window.mode_combo.setCurrentIndex(default_index)
        self.window.start_new_batch()
        previous_session_id = self.window.session.session_id

        self.assertTrue(self.window.mode_combo.isEnabled())
        new_api_index = self.window.mode_combo.findData(ModelMode.NEW_API.value)
        self.window.mode_combo.setCurrentIndex(new_api_index)
        with patch.object(
            QMessageBox,
            "question",
            return_value=QMessageBox.StandardButton.Yes,
        ):
            self.window.start_new_batch()

        self.assertNotEqual(previous_session_id, self.window.session.session_id)
        self.assertEqual(ModelMode.NEW_API, self.window.session.model_mode)
        self.assertEqual("Abandoned", self.store.get_session(previous_session_id).status)
        self.assertEqual(ModelMode.NEW_API.value, self.window.mode_combo.currentData())

    def test_copy_button_is_shorter_and_taller(self) -> None:
        self.assertEqual(156, self.window.copy_button.width())
        self.assertEqual(42, self.window.copy_button.height())

    def test_user_can_switch_to_another_prompt_sheet_for_new_batch(self) -> None:
        self.assertEqual(
            ["Cases", "long-text"],
            [
                self.window.prompt_sheet_combo.itemData(index)
                for index in range(self.window.prompt_sheet_combo.count())
            ],
        )
        self.window.start_new_batch()
        previous_session_id = self.window.session.session_id

        sheet_index = self.window.prompt_sheet_combo.findData("long-text")
        self.window.prompt_sheet_combo.setCurrentIndex(sheet_index)
        with patch.object(
            QMessageBox,
            "question",
            return_value=QMessageBox.StandardButton.Yes,
        ):
            self.window.start_new_batch()

        self.assertNotEqual(previous_session_id, self.window.session.session_id)
        self.assertEqual("long-text", self.window.session.prompt_sheet)
        self.assertEqual("Long01", self.window.current_item.case_name)
        self.assertEqual("这是长文本提示词", self.window.current_item.prompt)
        self.assertEqual("long-text", self.config.prompt_sheet)


if __name__ == "__main__":
    unittest.main()
