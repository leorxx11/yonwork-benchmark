from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from benchmark_companion.cases import CasesError, list_prompt_sheets, load_cases


class CasesTests(unittest.TestCase):
    def _workbook(self, rows: list[list[object]]) -> Path:
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        temp_dir = Path(temporary.name)
        path = temp_dir / "cases.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Cases"
        sheet.append(["CaseName", "Prompt", "Runs", "Enabled"])
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def test_loads_enabled_cases_and_runs(self) -> None:
        path = self._workbook(
            [["Case01", "你好", 2, True], ["Case02", "天气", 3, False]]
        )
        result = load_cases(path)
        self.assertEqual(2, len(result.cases))
        self.assertEqual(1, result.enabled_count)
        self.assertEqual(2, result.total_runs)
        self.assertEqual(["Case02 已禁用"], result.warnings)

    def test_rejects_duplicate_case_names(self) -> None:
        path = self._workbook(
            [["Case01", "你好", 1, True], ["Case01", "天气", 1, True]]
        )
        with self.assertRaisesRegex(CasesError, "重复"):
            load_cases(path)

    def test_lists_and_loads_all_compatible_prompt_sheets(self) -> None:
        path = self._workbook([["Case01", "你好", 1, True]])
        workbook = load_workbook(path)
        long_text = workbook.create_sheet("long-text")
        long_text.append(["CaseName", "Prompt", "Runs", "Enabled"])
        long_text.append(["Long01", "长文本提示词", 2, True])
        notes = workbook.create_sheet("Notes")
        notes.append(["Title", "Text"])
        workbook.save(path)
        workbook.close()

        self.assertEqual(["Cases", "long-text"], list_prompt_sheets(path))
        result = load_cases(path, sheet_name="long-text")
        self.assertEqual("Long01", result.cases[0].case_name)
        self.assertEqual(2, result.total_runs)


if __name__ == "__main__":
    unittest.main()
