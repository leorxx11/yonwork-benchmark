from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import CaseDefinition


EXPECTED_HEADERS = ("CaseName", "Prompt", "Runs", "Enabled")


class CasesError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class CaseLoadResult:
    cases: list[CaseDefinition]
    warnings: list[str]

    @property
    def enabled_count(self) -> int:
        return sum(1 for item in self.cases if item.enabled)

    @property
    def total_runs(self) -> int:
        return sum(item.runs for item in self.cases if item.enabled)


def list_prompt_sheets(workbook_path: Path) -> list[str]:
    """Return worksheets that use the supported prompt-table schema."""
    if not workbook_path.exists():
        raise CasesError(f"找不到工作簿：{workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        matches = [
            sheet.title
            for sheet in workbook.worksheets
            if tuple(sheet.cell(1, column).value for column in range(1, 5))
            == EXPECTED_HEADERS
        ]
    finally:
        workbook.close()

    if not matches:
        raise CasesError(
            "工作簿中没有兼容的提示词 Sheet；第一行前四列应为 "
            f"{EXPECTED_HEADERS}"
        )
    return matches


def _parse_enabled(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    if value is None:
        return False
    text = str(value).strip().casefold()
    if text in {"true", "yes", "y", "1", "是", "启用"}:
        return True
    if text in {"false", "no", "n", "0", "否", "禁用", ""}:
        return False
    raise CasesError(f"无法识别 Enabled 值：{value!r}")


def _parse_runs(value: Any, row_number: int, sheet_name: str) -> int:
    try:
        runs = int(value)
    except (TypeError, ValueError) as exc:
        raise CasesError(
            f"{sheet_name} 第 {row_number} 行的 Runs 不是整数：{value!r}"
        ) from exc
    if runs < 1:
        raise CasesError(f"{sheet_name} 第 {row_number} 行的 Runs 必须大于 0")
    return runs


def load_cases(workbook_path: Path, sheet_name: str = "Cases") -> CaseLoadResult:
    if not workbook_path.exists():
        raise CasesError(f"找不到工作簿：{workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise CasesError(f"工作簿缺少 {sheet_name} 工作表")
        sheet = workbook[sheet_name]
        headers = tuple(sheet.cell(1, column).value for column in range(1, 5))
        if headers != EXPECTED_HEADERS:
            raise CasesError(
                f"{sheet_name} 表头应为 {EXPECTED_HEADERS}，实际为 {headers}"
            )

        cases: list[CaseDefinition] = []
        warnings: list[str] = []
        seen_names: set[str] = set()
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, max_col=4, values_only=True), start=2
        ):
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            case_name = "" if values[0] is None else str(values[0]).strip()
            prompt = "" if values[1] is None else str(values[1]).strip()
            if not case_name:
                raise CasesError(f"{sheet_name} 第 {row_number} 行缺少 CaseName")
            if not prompt:
                raise CasesError(f"{sheet_name} 第 {row_number} 行缺少 Prompt")
            if case_name in seen_names:
                raise CasesError(f"{sheet_name} 中存在重复 CaseName：{case_name}")
            seen_names.add(case_name)

            runs = _parse_runs(values[2], row_number, sheet_name)
            enabled = _parse_enabled(values[3])
            if not enabled:
                warnings.append(f"{case_name} 已禁用")
            cases.append(CaseDefinition(case_name, prompt, runs, enabled))
    finally:
        workbook.close()

    if not cases:
        raise CasesError(f"{sheet_name} 工作表中没有可读取的数据")
    if not any(item.enabled for item in cases):
        raise CasesError(f"{sheet_name} 工作表中没有启用的 Case")
    return CaseLoadResult(cases=cases, warnings=warnings)
