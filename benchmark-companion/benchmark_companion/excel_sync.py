from __future__ import annotations

import os
import shutil
import tempfile
import zipfile
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.table import TableColumn

from .models import RunRecord, TokenStatus


BASE_HEADERS = (
    "Product",
    "Case",
    "Run",
    "Prompt",
    "StartTime",
    "EndTime",
    "DurationSeconds",
    "Model",
    "APICalls",
    "ErrorCalls",
    "InputTokens",
    "OutputTokens",
    "TotalTokens",
    "APIUseTime",
)

COMPANION_HEADERS = ("Status", "TokenStatus", "Note", "RecordId")
ALL_HEADERS = BASE_HEADERS + COMPANION_HEADERS
EXCEL_DATETIME_FORMAT = "yyyy/m/d hh:mm"
EXCEL_INTEGER_FORMAT = "0"


class WorkbookSyncError(RuntimeError):
    pass


@dataclass(frozen=True, slots=True)
class SyncOutcome:
    row_by_record: dict[str, int]
    backup_path: Path | None


def _excel_datetime(value: str | None) -> datetime | None:
    if not value:
        return None
    parsed = datetime.fromisoformat(value)
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone().replace(tzinfo=None)
    return parsed


def _contains_wps_cell_images(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as archive:
            names = {name.casefold() for name in archive.namelist()}
            risky = {
                "xl/cellimages.xml",
                "xl/_rels/cellimages.xml.rels",
                "xl/woinfos.xml",
            }
            return bool(names & risky)
    except zipfile.BadZipFile as exc:
        raise WorkbookSyncError(f"工作簿不是有效的 XLSX 文件：{path}") from exc


class ExcelResultsSync:
    def __init__(self, workbook_path: Path, backup_dir: Path):
        self.workbook_path = workbook_path
        self.backup_dir = backup_dir

    def sync_records(self, records: Sequence[RunRecord]) -> SyncOutcome:
        if not records:
            return SyncOutcome(row_by_record={}, backup_path=None)
        if not self.workbook_path.exists():
            raise WorkbookSyncError(f"找不到工作簿：{self.workbook_path}")
        if any(Path(record.workbook_path).resolve() != self.workbook_path.resolve() for record in records):
            raise WorkbookSyncError("一次同步只能写入同一个工作簿")
        if _contains_wps_cell_images(self.workbook_path):
            raise WorkbookSyncError(
                "检测到 WPS 单元格图片部件，为避免损坏，已停止自动写入"
            )

        keep_vba = self.workbook_path.suffix.casefold() == ".xlsm"
        try:
            workbook = load_workbook(
                self.workbook_path,
                read_only=False,
                data_only=False,
                keep_vba=keep_vba,
                keep_links=True,
            )
        except PermissionError as exc:
            raise WorkbookSyncError("工作簿正在被占用，结果已保留在本地") from exc
        except Exception as exc:
            raise WorkbookSyncError(f"无法打开工作簿：{exc}") from exc

        temp_path: Path | None = None
        backup_path: Path | None = None
        try:
            if "Results" not in workbook.sheetnames:
                raise WorkbookSyncError("工作簿缺少 Results 工作表")
            sheet = workbook["Results"]
            self._ensure_headers(sheet)

            last_data_row = self._last_populated_row(sheet)
            existing_rows = {
                str(sheet.cell(row, 18).value): row
                for row in range(2, last_data_row + 1)
                if sheet.cell(row, 18).value not in (None, "")
            }

            row_by_record: dict[str, int] = {}
            for record in records:
                row_number = existing_rows.get(record.record_id)
                if row_number is None:
                    row_number = last_data_row + 1
                    self._copy_row_style(sheet, last_data_row, row_number)
                    last_data_row = row_number
                    existing_rows[record.record_id] = row_number
                self._write_record(sheet, row_number, record)
                row_by_record[record.record_id] = row_number

            had_results_table = self._extend_results_table(sheet, last_data_row)
            self._apply_number_formats(sheet, row_by_record.values())

            file_descriptor, raw_temp = tempfile.mkstemp(
                prefix=f".{self.workbook_path.stem}.companion-",
                suffix=self.workbook_path.suffix,
                dir=self.workbook_path.parent,
            )
            os.close(file_descriptor)
            temp_path = Path(raw_temp)
            workbook.save(temp_path)
            workbook.close()

            self._verify_saved_file(temp_path, row_by_record, require_table=had_results_table)
            backup_path = self._create_daily_backup()
            try:
                os.replace(temp_path, self.workbook_path)
            except PermissionError as exc:
                raise WorkbookSyncError("工作簿正在被占用，结果已保留在本地") from exc
            temp_path = None
            return SyncOutcome(row_by_record=row_by_record, backup_path=backup_path)
        except WorkbookSyncError:
            raise
        except PermissionError as exc:
            raise WorkbookSyncError("工作簿正在被占用，结果已保留在本地") from exc
        except Exception as exc:
            raise WorkbookSyncError(f"同步 Excel 失败：{exc}") from exc
        finally:
            try:
                workbook.close()
            except Exception:
                pass
            if temp_path is not None:
                try:
                    temp_path.unlink(missing_ok=True)
                except OSError:
                    pass

    @staticmethod
    def _ensure_headers(sheet) -> None:
        actual_base = tuple(sheet.cell(1, column).value for column in range(1, 15))
        if actual_base != BASE_HEADERS:
            raise WorkbookSyncError(
                f"Results 前 14 列表头不匹配，实际为：{actual_base}"
            )
        header_style_source = sheet.cell(1, 14)
        new_widths = {15: 12, 16: 16, 17: 34, 18: 40}
        for column, header in enumerate(COMPANION_HEADERS, start=15):
            cell = sheet.cell(1, column)
            if cell.value not in (None, "", header):
                raise WorkbookSyncError(
                    f"Results 第 {column} 列已被占用：{cell.value!r}"
                )
            cell.value = header
            if header_style_source.has_style:
                cell._style = copy(header_style_source._style)
            if header_style_source.alignment:
                cell.alignment = copy(header_style_source.alignment)
            if header_style_source.font:
                cell.font = copy(header_style_source.font)
            if header_style_source.fill:
                cell.fill = copy(header_style_source.fill)
            if header_style_source.border:
                cell.border = copy(header_style_source.border)
            dimension = sheet.column_dimensions[get_column_letter(column)]
            if dimension.width is None or dimension.width < new_widths[column]:
                dimension.width = new_widths[column]

    @staticmethod
    def _last_populated_row(sheet) -> int:
        last = 1
        for row in range(2, sheet.max_row + 1):
            if any(sheet.cell(row, column).value not in (None, "") for column in range(1, 19)):
                last = row
        return last

    @staticmethod
    def _copy_row_style(sheet, source_row: int, target_row: int) -> None:
        if source_row < 2:
            source_row = 2
        for column in range(1, 19):
            source_column = column if column <= 14 else 14
            source = sheet.cell(source_row, source_column)
            target = sheet.cell(target_row, column)
            if source.has_style:
                target._style = copy(source._style)
            target.alignment = copy(source.alignment)
            target.protection = copy(source.protection)
        if source_row in sheet.row_dimensions:
            sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height

    @staticmethod
    def _write_record(sheet, row: int, record: RunRecord) -> None:
        has_token_data = record.token_status is TokenStatus.SUCCESS
        values = (
            record.product,
            record.case_name,
            record.run_no,
            record.prompt,
            _excel_datetime(record.start_time),
            _excel_datetime(record.end_time),
            record.duration_seconds,
            record.model,
            record.api_calls if has_token_data else None,
            record.error_calls if has_token_data else None,
            record.input_tokens if has_token_data else None,
            record.output_tokens if has_token_data else None,
            record.total_tokens if has_token_data else None,
            record.api_use_time if has_token_data else None,
            record.status.value,
            record.token_status.value,
            record.note,
            record.record_id,
        )
        for column, value in enumerate(values, start=1):
            sheet.cell(row, column).value = value

    @staticmethod
    def _extend_results_table(sheet, last_data_row: int) -> bool:
        matching_table = None
        for table in sheet.tables.values():
            min_col, min_row, max_col, _ = range_boundaries(table.ref)
            if min_row == 1 and min_col == 1 and max_col >= 14:
                matching_table = table
                break
        if matching_table is None:
            return False

        columns = list(matching_table.tableColumns)
        if columns and len(columns) < len(ALL_HEADERS):
            next_id = max(column.id for column in columns) + 1
            for header in ALL_HEADERS[len(columns) :]:
                matching_table.tableColumns.append(TableColumn(id=next_id, name=header))
                next_id += 1
        elif not columns:
            for index, header in enumerate(ALL_HEADERS, start=1):
                matching_table.tableColumns.append(TableColumn(id=index, name=header))
        elif len(columns) > len(ALL_HEADERS):
            raise WorkbookSyncError("Results 表格包含未识别的额外列")

        matching_table.ref = f"A1:R{max(last_data_row, 2)}"
        if matching_table.autoFilter is not None:
            matching_table.autoFilter.ref = matching_table.ref
        return True

    @staticmethod
    def _apply_number_formats(sheet, rows) -> None:
        for row in rows:
            sheet.cell(row, 5).number_format = EXCEL_DATETIME_FORMAT
            sheet.cell(row, 6).number_format = EXCEL_DATETIME_FORMAT
            sheet.cell(row, 7).number_format = "0.000"
            for column in range(9, 14):
                sheet.cell(row, column).number_format = EXCEL_INTEGER_FORMAT
            sheet.cell(row, 14).number_format = "0.000"

    def _create_daily_backup(self) -> Path:
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        day = datetime.now().strftime("%Y%m%d")
        backup_path = self.backup_dir / f"{self.workbook_path.stem}.{day}{self.workbook_path.suffix}"
        if not backup_path.exists():
            shutil.copy2(self.workbook_path, backup_path)
        return backup_path

    @staticmethod
    def _verify_saved_file(
        path: Path, row_by_record: dict[str, int], *, require_table: bool
    ) -> None:
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            if "Results" not in workbook.sheetnames:
                raise WorkbookSyncError("保存校验失败：Results 工作表丢失")
            sheet = workbook["Results"]
            headers = tuple(sheet.cell(1, column).value for column in range(1, 19))
            if headers != ALL_HEADERS:
                raise WorkbookSyncError("保存校验失败：Results 表头不完整")
            for record_id, row in row_by_record.items():
                if str(sheet.cell(row, 18).value) != record_id:
                    raise WorkbookSyncError("保存校验失败：RecordId 未正确写入")
            if require_table:
                table_refs = [table.ref for table in sheet.tables.values()]
                if not any(range_boundaries(ref)[2] >= 18 for ref in table_refs):
                    raise WorkbookSyncError("保存校验失败：Results 表格未扩展到 R 列")
        finally:
            workbook.close()
