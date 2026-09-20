from __future__ import annotations

import json
import sqlite3
from collections import Counter
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterator

from .models import USAGE_SOURCES, JsonObject, RunRecord, Verdict


SCHEMA = """
CREATE TABLE IF NOT EXISTS runs (
    benchmark_id TEXT PRIMARY KEY,
    batch_id TEXT NOT NULL,
    position INTEGER NOT NULL,
    case_name TEXT NOT NULL,
    run_no INTEGER NOT NULL,
    prompt TEXT NOT NULL,
    session_key TEXT NOT NULL,
    verdict TEXT NOT NULL,
    run_id TEXT,
    answer TEXT,
    started_at TEXT,
    ended_at TEXT,
    duration_seconds REAL,
    first_delta_seconds REAL,
    terminated_by TEXT,
    stop_reason TEXT,
    tool_call_count INTEGER,
    input_tokens INTEGER,
    output_tokens INTEGER,
    total_tokens INTEGER,
    cost_usd REAL,
    usage_match TEXT,
    model TEXT,
    usage_source TEXT,
    requested_model TEXT,
    api_calls INTEGER,
    error_calls INTEGER,
    note TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS checks (
    benchmark_id TEXT NOT NULL,
    layer TEXT NOT NULL,
    name TEXT NOT NULL,
    verdict TEXT,
    detail TEXT NOT NULL DEFAULT '',
    PRIMARY KEY (benchmark_id, layer, name),
    FOREIGN KEY (benchmark_id) REFERENCES runs(benchmark_id)
);

CREATE INDEX IF NOT EXISTS idx_runs_verdict ON runs(batch_id, verdict);
CREATE INDEX IF NOT EXISTS idx_checks_verdict ON checks(verdict);
"""

RESULT_COLUMNS = (
    ("benchmark_id", "BenchmarkId"),
    ("case_name", "Case"),
    ("run_no", "Run"),
    ("verdict", "Verdict"),
    ("started_at", "StartTime"),
    ("ended_at", "EndTime"),
    ("duration_seconds", "DurationSeconds"),
    ("first_delta_seconds", "FirstDeltaSeconds"),
    ("terminated_by", "TerminatedBy"),
    ("stop_reason", "StopReason"),
    ("requested_model", "RequestedModel"),
    ("model", "ActualModel"),
    ("api_calls", "APICalls"),
    ("error_calls", "ErrorCalls"),
    ("input_tokens", "InputTokens"),
    ("output_tokens", "OutputTokens"),
    ("total_tokens", "TotalTokens"),
    ("cost_usd", "CostUsd"),
    ("usage_source", "UsageSource"),
    ("usage_match", "UsageMatch"),
    ("tool_call_count", "ToolCalls"),
    ("session_key", "SessionKey"),
    ("run_id", "RunId"),
    ("note", "Note"),
)

_ANSWER_PREVIEW_CHARS = 2000


class ReportError(RuntimeError):
    pass


def append_jsonl(path: Path, record: RunRecord) -> None:
    """一轮一行，立刻落盘。

    整批结束才写就等于回到 PAD 时代：跑到一半崩了什么都不剩。
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(record.to_json(), ensure_ascii=False) + "\n")
        handle.flush()


def iter_jsonl(path: Path) -> Iterator[JsonObject]:
    if not path.is_file():
        raise ReportError(f"找不到结果文件：{path}")
    with path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                value = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ReportError(f"{path} 第 {line_number} 行不是合法 JSON") from exc
            if not isinstance(value, dict):
                raise ReportError(f"{path} 第 {line_number} 行不是 JSON 对象")
            yield value


@contextmanager
def _connection(database_path: Path) -> Iterator[sqlite3.Connection]:
    database_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(database_path, timeout=15)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        yield connection
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def usage_samples_of(record: JsonObject) -> list[JsonObject]:
    """取一条记录里的全部用量来源。

    兼容早期格式：那时 `usage` 是单个对象，只有端上一路。
    """
    samples = record.get("usage_samples")
    if isinstance(samples, list):
        return [item for item in samples if isinstance(item, dict)]
    legacy = record.get("usage")
    if isinstance(legacy, dict):
        return [{"source": "device-api", **legacy}]
    return []


def primary_usage(record: JsonObject) -> JsonObject | None:
    samples = usage_samples_of(record)
    for source in USAGE_SOURCES:
        for sample in samples:
            if sample.get("source") == source:
                return sample
    return None


def _get(record: JsonObject, section: str, key: str) -> Any:
    block = record.get(section)
    return block.get(key) if isinstance(block, dict) else None


def _row(record: JsonObject) -> JsonObject:
    turn = record.get("turn") if isinstance(record.get("turn"), dict) else {}
    tool_calls = turn.get("tool_calls") or []
    answer = turn.get("answer")
    usage = primary_usage(record) or {}
    return {
        "benchmark_id": record.get("benchmark_id"),
        "batch_id": record.get("batch_id"),
        "position": record.get("position", 0),
        "case_name": record.get("case_name", ""),
        "run_no": record.get("run_no", 0),
        "prompt": record.get("prompt", ""),
        "session_key": record.get("session_key", ""),
        "verdict": record.get("verdict", Verdict.ERROR.value),
        "run_id": turn.get("run_id"),
        "answer": answer[:_ANSWER_PREVIEW_CHARS] if isinstance(answer, str) else None,
        "started_at": turn.get("started_at"),
        "ended_at": turn.get("ended_at"),
        "duration_seconds": turn.get("duration_seconds"),
        "first_delta_seconds": turn.get("first_delta_seconds"),
        "terminated_by": turn.get("terminated_by"),
        "stop_reason": turn.get("stop_reason"),
        "tool_call_count": len(tool_calls) if isinstance(tool_calls, list) else None,
        "input_tokens": usage.get("input_tokens"),
        "output_tokens": usage.get("output_tokens"),
        "total_tokens": usage.get("total_tokens"),
        "cost_usd": usage.get("cost_usd"),
        "usage_match": usage.get("match"),
        "usage_source": usage.get("source"),
        "model": usage.get("model"),
        "requested_model": turn.get("requested_model"),
        "api_calls": _get(record, "log_stats", "api_calls"),
        "error_calls": _get(record, "log_stats", "error_calls"),
        "note": record.get("note", "") or "",
        "created_at": record.get("created_at", ""),
    }


def build_database(jsonl_path: Path, database_path: Path) -> int:
    """JSONL → SQLite。重跑同一个 BenchmarkId 会覆盖旧行。"""

    rows: list[JsonObject] = []
    checks: list[tuple[Any, ...]] = []
    for record in iter_jsonl(jsonl_path):
        rows.append(_row(record))
        for check in record.get("checks") or []:
            if not isinstance(check, dict):
                continue
            checks.append(
                (
                    record.get("benchmark_id"),
                    check.get("layer"),
                    check.get("name"),
                    check.get("verdict"),
                    check.get("detail", ""),
                )
            )

    with _connection(database_path) as connection:
        connection.executescript(SCHEMA)
        columns = list(rows[0].keys()) if rows else []
        if rows:
            placeholders = ", ".join(f":{name}" for name in columns)
            connection.executemany(
                f"INSERT OR REPLACE INTO runs ({', '.join(columns)}) VALUES ({placeholders})",
                rows,
            )
        connection.executemany(
            "INSERT OR REPLACE INTO checks (benchmark_id, layer, name, verdict, detail)"
            " VALUES (?, ?, ?, ?, ?)",
            checks,
        )
    return len(rows)


@dataclass(frozen=True, slots=True)
class Summary:
    total: int
    by_verdict: dict[str, int]

    @property
    def worst_verdict(self) -> Verdict:
        from .models import worst

        present = [
            Verdict(name) for name, count in self.by_verdict.items() if count > 0
        ]
        return worst(present)

    def as_text(self) -> str:
        parts = [
            f"{verdict.value}={self.by_verdict.get(verdict.value, 0)}"
            for verdict in Verdict
        ]
        return f"共 {self.total} 轮：" + "、".join(parts)


def summarize(jsonl_path: Path) -> Summary:
    counter: Counter[str] = Counter()
    total = 0
    for record in iter_jsonl(jsonl_path):
        total += 1
        counter[str(record.get("verdict", Verdict.ERROR.value))] += 1
    return Summary(total=total, by_verdict=dict(counter))


def export_xlsx(database_path: Path, xlsx_path: Path) -> Path:
    """SQLite → xlsx。Results 一行一轮，Checks 一行一条断言。"""

    from openpyxl import Workbook

    with _connection(database_path) as connection:
        runs = connection.execute(
            "SELECT * FROM runs ORDER BY batch_id, position"
        ).fetchall()
        checks = connection.execute(
            "SELECT * FROM checks ORDER BY benchmark_id, layer, name"
        ).fetchall()

    workbook = Workbook()
    results = workbook.active
    results.title = "Results"
    results.append([title for _, title in RESULT_COLUMNS])
    for row in runs:
        results.append([row[column] for column, _ in RESULT_COLUMNS])

    sheet = workbook.create_sheet("Checks")
    sheet.append(["BenchmarkId", "Layer", "Check", "Verdict", "Detail"])
    for row in checks:
        sheet.append(
            [
                row["benchmark_id"],
                row["layer"],
                row["name"],
                row["verdict"] or "Skipped",
                row["detail"],
            ]
        )

    counter: Counter[str] = Counter(str(row["verdict"]) for row in runs)
    sheet = workbook.create_sheet("Summary")
    sheet.append(["Verdict", "Count"])
    for verdict in Verdict:
        sheet.append([verdict.value, counter.get(verdict.value, 0)])
    sheet.append(["Total", len(runs)])

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)
    workbook.close()
    return xlsx_path
