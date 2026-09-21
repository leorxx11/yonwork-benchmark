from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import CaseDefinition, Expectations


# 前四列沿用 Benchmark Companion 的表结构，yonwork_benchmark.xlsx 现成可用。
REQUIRED_HEADERS = ("CaseName", "Prompt", "Runs", "Enabled")

# 断言参数列，全部可选；表里没有就用 Expectations 的默认值。
# 多值列用 | 分隔（提示词里逗号太常见，用逗号会误切）。
OPTIONAL_HEADERS = (
    "Expect",
    "Forbid",
    "MinLength",
    "JsonParsable",
    "MaxSeconds",
    "MaxTotalTokens",
    "MaxInputTokens",
)
MULTI_VALUE_SEPARATOR = "|"


class CasesError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class CaseLoadResult:
    cases: list[CaseDefinition]
    warnings: list[str]

    @property
    def enabled_count(self) -> int:
        return sum(1 for item in self.cases if item.enabled)

    @property
    def total_runs(self) -> int:
        return sum(item.runs for item in self.cases if item.enabled)


def list_prompt_sheets(workbook_path: Path) -> list[str]:
    """列出用了提示词表结构的工作表。"""
    if not workbook_path.exists():
        raise CasesError(f"找不到工作簿：{workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        matches = [
            sheet.title
            for sheet in workbook.worksheets
            if tuple(sheet.cell(1, column).value for column in range(1, 5))
            == REQUIRED_HEADERS
        ]
    finally:
        workbook.close()

    if not matches:
        raise CasesError(
            f"工作簿中没有兼容的提示词 Sheet；第一行前四列应为 {REQUIRED_HEADERS}"
        )
    return matches


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _parse_enabled(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    text = _text(value).casefold()
    if text in {"true", "yes", "y", "1", "是", "启用"}:
        return True
    if text in {"false", "no", "n", "0", "否", "禁用", ""}:
        return False
    raise CasesError(f"无法识别 Enabled 值：{value!r}")


def _parse_runs(value: Any, row_number: int, sheet_name: str) -> int:
    try:
        runs = int(value)
    except (TypeError, ValueError) as exc:
        raise CasesError(
            f"{sheet_name} 第 {row_number} 行的 Runs 不是整数：{value!r}"
        ) from exc
    if runs < 1:
        raise CasesError(f"{sheet_name} 第 {row_number} 行的 Runs 必须大于 0")
    return runs


def _parse_keywords(value: Any) -> tuple[str, ...]:
    text = _text(value)
    if not text:
        return ()
    return tuple(
        item.strip() for item in text.split(MULTI_VALUE_SEPARATOR) if item.strip()
    )


def _parse_number(
    value: Any, column: str, row_number: int, sheet_name: str, *, cast: type
) -> Any:
    text = _text(value)
    if not text:
        return None
    try:
        parsed = cast(float(text))
    except (TypeError, ValueError) as exc:
        raise CasesError(
            f"{sheet_name} 第 {row_number} 行的 {column} 不是数字：{value!r}"
        ) from exc
    if parsed <= 0:
        raise CasesError(f"{sheet_name} 第 {row_number} 行的 {column} 必须大于 0")
    return parsed


def _parse_expectations(
    row: dict[str, Any], row_number: int, sheet_name: str
) -> Expectations:
    min_length = _parse_number(
        row.get("MinLength"), "MinLength", row_number, sheet_name, cast=int
    )
    return Expectations(
        expect_keywords=_parse_keywords(row.get("Expect")),
        forbid_keywords=_parse_keywords(row.get("Forbid")),
        min_length=min_length if min_length is not None else 1,
        json_parsable=_parse_enabled(row.get("JsonParsable")),
        max_seconds=_parse_number(
            row.get("MaxSeconds"), "MaxSeconds", row_number, sheet_name, cast=float
        ),
        max_total_tokens=_parse_number(
            row.get("MaxTotalTokens"),
            "MaxTotalTokens",
            row_number,
            sheet_name,
            cast=int,
        ),
        max_input_tokens=_parse_number(
            row.get("MaxInputTokens"),
            "MaxInputTokens",
            row_number,
            sheet_name,
            cast=int,
        ),
    )


def load_cases(workbook_path: Path, sheet_name: str = "Cases") -> CaseLoadResult:
    if not workbook_path.exists():
        raise CasesError(f"找不到工作簿：{workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise CasesError(f"工作簿缺少 {sheet_name} 工作表")
        sheet = workbook[sheet_name]

        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [_text(value) for value in header_row]
        if tuple(headers[:4]) != REQUIRED_HEADERS:
            raise CasesError(
                f"{sheet_name} 表头前四列应为 {REQUIRED_HEADERS}，实际为 {tuple(headers[:4])}"
            )
        unknown = [
            name
            for name in headers[4:]
            if name and name not in OPTIONAL_HEADERS
        ]

        cases: list[CaseDefinition] = []
        warnings: list[str] = [f"忽略未知列：{name}" for name in unknown]
        seen_names: set[str] = set()
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True), start=2
        ):
            if all(value is None or _text(value) == "" for value in values):
                continue
            row = dict(zip(headers, values))
            case_name = _text(values[0])
            prompt = _text(values[1])
            if not case_name:
                raise CasesError(f"{sheet_name} 第 {row_number} 行缺少 CaseName")
            if not prompt:
                raise CasesError(f"{sheet_name} 第 {row_number} 行缺少 Prompt")
            if case_name in seen_names:
                raise CasesError(f"{sheet_name} 中存在重复 CaseName：{case_name}")
            seen_names.add(case_name)

            enabled = _parse_enabled(values[3])
            if not enabled:
                warnings.append(f"{case_name} 已禁用")
            cases.append(
                CaseDefinition(
                    case_name=case_name,
                    prompt=prompt,
                    runs=_parse_runs(values[2], row_number, sheet_name),
                    enabled=enabled,
                    expectations=_parse_expectations(row, row_number, sheet_name),
                )
            )
    finally:
        workbook.close()

    if not cases:
        raise CasesError(f"{sheet_name} 工作表中没有可读取的数据")
    if not any(item.enabled for item in cases):
        raise CasesError(f"{sheet_name} 工作表中没有启用的 Case")
    return CaseLoadResult(cases=cases, warnings=warnings)
